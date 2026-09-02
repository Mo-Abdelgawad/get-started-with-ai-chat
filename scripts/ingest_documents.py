from __future__ import annotations

import csv
import json
import os
import sys
import tempfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional

from azure.core.credentials_async import AsyncTokenCredential
from azure.identity.aio import AzureDeveloperCliCredential
from openai import AsyncAzureOpenAI

from src.api.search_index_manager import SearchIndexManager


SUPPORTED_SECTORS = {
    "healthcare",
    "education",
    "finance",
    "driving_licensing",
    "energy",
    "retail",
    "public_sector",
}

SUPPORTED_REGIONS = {
    "Aleppo",
    "Damascus",
    "Latakia",
    "national",
    "regional",
}

SEARCH_METADATA_FIELDS = [
    "sector",
    "region",
    "document_type",
    "source_owner",
    "language",
    "tags",
    "confidential_level",
]


@dataclass
class DocumentMetadata:
    document_id: str
    title: str
    source_file: str
    sector: str = "unknown"
    region: str = "unknown"
    document_type: str = "unknown"
    language: str = "en"
    source_owner: str = "unknown"
    publish_date: Optional[str] = None
    tags: List[str] = field(default_factory=list)
    confidential_level: str = "internal"
    page_count: int = 1
    ingestion_status: str = "pending"

    def to_dict(self) -> Dict[str, Any]:
        payload = asdict(self)
        payload["tags"] = list(self.tags)
        return payload


def infer_sector_from_name(file_name: str, override: Optional[str] = None) -> str:
    if override:
        return override
    candidate = file_name.lower()
    mapping = {
        "healthcare": ["healthcare", "hospital", "clinic", "medical"],
        "education": ["education", "school", "university", "training"],
        "finance": ["finance", "bank", "insurance", "investment"],
        "driving_licensing": ["driving", "license", "licence", "driver"],
        "energy": ["energy", "electricity", "oil", "gas"],
        "retail": ["retail", "market", "commerce", "sales"],
        "public_sector": ["government", "public", "municipality", "policy"],
    }
    for sector, tokens in mapping.items():
        if any(token in candidate for token in tokens):
            return sector
    return "unknown"


def infer_region_from_name(file_name: str, override: Optional[str] = None) -> str:
    if override:
        return override
    candidate = file_name.lower()
    mapping = {
        "Aleppo": ["aleppo"],
        "Damascus": ["damascus"],
        "Latakia": ["latakia"],
        "national": ["national", "countrywide"],
        "regional": ["regional"],
    }
    for region, tokens in mapping.items():
        if any(token in candidate for token in tokens):
            return region
    return "unknown"


def infer_document_type(file_name: str, override: Optional[str] = None) -> str:
    if override:
        return override
    suffix = Path(file_name).suffix.lower().lstrip(".")
    mapping = {
        "pdf": "pdf",
        "md": "markdown",
        "txt": "text",
        "csv": "csv",
        "docx": "word",
        "doc": "word",
        "xlsx": "excel",
        "xls": "excel",
        "pptx": "powerpoint",
    }
    return mapping.get(suffix, suffix or "unknown")


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 120) -> List[str]:
    if not text:
        return []

    normalized = " ".join(text.split())
    if len(normalized) <= chunk_size:
        return [normalized]

    chunks: List[str] = []
    start = 0
    while start < len(normalized):
        end = min(start + chunk_size, len(normalized))
        chunk = normalized[start:end]
        chunks.append(chunk)
        if end >= len(normalized):
            break
        start = max(end - overlap, start + 1)
    return chunks


def extract_pdf_text(file_path: Path) -> str:
    try:
        import fitz
    except ModuleNotFoundError as exc:  # pragma: no cover
        raise RuntimeError("PyMuPDF is required for PDF ingestion. Install it with: pip install pymupdf") from exc

    doc = fitz.open(file_path)
    try:
        pages: List[str] = []
        for page in doc:
            pages.append(page.get_text("text"))
        return "\n\n".join(page for page in pages if page)
    finally:
        doc.close()


def extract_docx_text(file_path: Path) -> str:
    try:
        from docx import Document
    except ModuleNotFoundError as exc:  # pragma: no cover
        raise RuntimeError("python-docx is required for DOCX ingestion. Install it with: pip install python-docx") from exc

    document = Document(str(file_path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    return "\n".join(paragraphs)


def extract_xlsx_text(file_path: Path) -> str:
    try:
        import openpyxl
    except ModuleNotFoundError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required for XLSX ingestion. Install it with: pip install openpyxl") from exc

    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    rows: List[str] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell) if cell is not None else "" for cell in row]
                rows.append(" | ".join(cells))
        return "\n".join(rows)
    finally:
        workbook.close()


def extract_csv_text(file_path: Path) -> str:
    rows: List[str] = []
    csv.field_size_limit(sys.maxsize)
    with file_path.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
        reader = csv.reader(handle)
        for row in reader:
            rows.append(" | ".join(row))
    return "\n".join(rows)


def extract_markdown_text(file_path: Path) -> str:
    return file_path.read_text(encoding="utf-8", errors="ignore")


def extract_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_text(file_path)
    if suffix == ".docx":
        return extract_docx_text(file_path)
    if suffix in {".xlsx", ".xls"}:
        return extract_xlsx_text(file_path)
    if suffix == ".csv":
        return extract_csv_text(file_path)
    if suffix in {".md", ".txt"}:
        return extract_markdown_text(file_path)
    raise ValueError(f"Unsupported file type for ingestion: {file_path}")


def build_metadata(file_name: str, *, sector_override: Optional[str] = None, region_override: Optional[str] = None, source_owner: Optional[str] = None) -> DocumentMetadata:
    document_id = file_name
    title = Path(file_name).stem
    metadata = DocumentMetadata(
        document_id=document_id,
        title=title,
        source_file=file_name,
        sector=infer_sector_from_name(file_name, sector_override),
        region=infer_region_from_name(file_name, region_override),
        document_type=infer_document_type(file_name),
        source_owner=source_owner or "unknown",
    )
    return metadata


def ingest_documents(
    source_directory: str,
    *,
    chunk_size: int = 800,
    chunk_overlap: int = 120,
    sector_override: Optional[str] = None,
    region_override: Optional[str] = None,
    source_owner: Optional[str] = None,
    on_chunk: Optional[Callable[[DocumentMetadata, str, int], None]] = None,
) -> List[Dict[str, Any]]:
    """Generic ingestion skeleton for multi-sector document sources.

    Each chunk is returned as a dictionary and can be sent to Azure AI Search or any custom indexer.
    """
    ingested: List[Dict[str, Any]] = []
    for file_path in sorted(Path(source_directory).iterdir()):
        if not file_path.is_file():
            continue
        if file_path.name.startswith("."):
            continue
        if file_path.suffix.lower() not in {".pdf", ".md", ".txt", ".csv", ".docx", ".xlsx", ".xls"}:
            continue

        metadata = build_metadata(
            file_path.name,
            sector_override=sector_override,
            region_override=region_override,
            source_owner=source_owner,
        )

        text = extract_text(file_path)
        metadata.page_count = max(1, text.count("\n") // 35 + 1)
        chunks = chunk_text(text, chunk_size=chunk_size, overlap=chunk_overlap)

        for index, chunk in enumerate(chunks):
            payload = {
                "document_id": metadata.document_id,
                "title": metadata.title,
                "source_file": metadata.source_file,
                "sector": metadata.sector,
                "region": metadata.region,
                "document_type": metadata.document_type,
                "language": metadata.language,
                "source_owner": metadata.source_owner,
                "publish_date": metadata.publish_date,
                "tags": metadata.tags,
                "confidential_level": metadata.confidential_level,
                "page_count": metadata.page_count,
                "ingestion_status": "processed",
                "chunk_index": index,
                "chunk_text": chunk,
            }
            ingested.append(payload)
            if on_chunk is not None:
                on_chunk(metadata, chunk, index)

    return ingested


def write_search_csv(records: Iterable[Dict[str, Any]], output_path: Path) -> None:
    """Convert chunk records to the Azure Search CSV schema used by the app."""
    rows = list(records)
    if not rows:
        raise ValueError("No records available for Azure Search upload.")

    fieldnames = ["token", "source_file", "embedding"] + SEARCH_METADATA_FIELDS
    with output_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({
                "token": row["chunk_text"],
                "source_file": row["source_file"],
                "embedding": json.dumps(row["embedding"]),
                "sector": row.get("sector", "unknown"),
                "region": row.get("region", "unknown"),
                "document_type": row.get("document_type", "unknown"),
                "source_owner": row.get("source_owner", "unknown"),
                "language": row.get("language", "en"),
                "tags": ", ".join(row.get("tags", []) or []),
                "confidential_level": row.get("confidential_level", "internal"),
            })


async def upload_records_to_search(
    records: Iterable[Dict[str, Any]],
    *,
    endpoint: Optional[str] = None,
    index_name: Optional[str] = None,
    dimensions: Optional[int] = None,
    model: Optional[str] = None,
    credential: Optional[AsyncTokenCredential] = None,
    output_path: Optional[Path] = None,
) -> None:
    """Embed and upload metadata-tagged chunks to Azure AI Search."""
    rows = list(records)
    if not rows:
        raise ValueError("No records available for Azure Search upload.")

    if endpoint is None:
        endpoint = os.environ.get("AZURE_AI_SEARCH_ENDPOINT")
    if index_name is None:
        index_name = os.environ.get("AZURE_AI_SEARCH_INDEX_NAME", "qarar-citations-v1")
    if model is None:
        model = os.environ.get("AZURE_AI_EMBED_DEPLOYMENT_NAME", "text-embedding-3-small")
    if dimensions is None:
        dimensions = int(os.environ.get("AZURE_AI_EMBED_DIMENSIONS", "1536"))
    if credential is None:
        credential = AzureDeveloperCliCredential()

    if endpoint is None:
        raise ValueError("AZURE_AI_SEARCH_ENDPOINT is required for upload.")

    openai_endpoint = os.environ.get("AZURE_OPENAI_ENDPOINT")
    if openai_endpoint is None:
        raise ValueError("AZURE_OPENAI_ENDPOINT is required to create embeddings.")

    async def get_openai_token() -> str:
        return (await credential.get_token(
            "https://cognitiveservices.azure.com/.default"
        )).token

    embed = AsyncAzureOpenAI(
        azure_endpoint=openai_endpoint,
        api_version=os.getenv("AZURE_OPENAI_API_VERSION", "2024-10-21"),
        azure_ad_token_provider=get_openai_token,
    )

    try:
        manager = SearchIndexManager(
            endpoint=endpoint,
            credential=credential,
            index_name=index_name,
            dimensions=dimensions,
            model=model,
            embeddings_client=embed,
        )
        temp_csv = output_path or Path(tempfile.gettempdir()) / "qarar_ingest_upload.csv"

        metadata_payload = []
        for row in rows:
            metadata_payload.append({
                "chunk_text": row["chunk_text"],
                "source_file": row["source_file"],
                "sector": row.get("sector", "unknown"),
                "region": row.get("region", "unknown"),
                "document_type": row.get("document_type", "unknown"),
                "source_owner": row.get("source_owner", "unknown"),
                "language": row.get("language", "en"),
                "tags": row.get("tags", []),
                "confidential_level": row.get("confidential_level", "internal"),
            })

        async def _embed_and_write() -> None:
            csv_rows = []
            for metadata_row in metadata_payload:
                embedding_response = await embed.embeddings.create(
                    input=[metadata_row["chunk_text"]],
                    model=model,
                )
                embedding = embedding_response.data[0].embedding
                csv_rows.append({
                    "token": metadata_row["chunk_text"],
                    "source_file": metadata_row["source_file"],
                    "embedding": json.dumps(embedding),
                    "sector": metadata_row["sector"],
                    "region": metadata_row["region"],
                    "document_type": metadata_row["document_type"],
                    "source_owner": metadata_row["source_owner"],
                    "language": metadata_row["language"],
                    "tags": ", ".join(metadata_row["tags"] or []),
                    "confidential_level": metadata_row["confidential_level"],
                })

            with temp_csv.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=["token", "source_file", "embedding", *SEARCH_METADATA_FIELDS])
                writer.writeheader()
                writer.writerows(csv_rows)

        await _embed_and_write()
        await manager.ensure_index_created(
            vector_index_dimensions=dimensions,
            metadata_fields=SEARCH_METADATA_FIELDS,
        )
        await manager.upload_documents(str(temp_csv))
        print(f"Uploaded {len(rows)} records to Azure AI Search index '{index_name}'.")
    finally:
        await manager.close()
        await embed.close()


async def _example_upload() -> None:
    sample_dir = Path(__file__).resolve().parent.parent / "data"
    rows = ingest_documents(str(sample_dir))
    await upload_records_to_search(rows)


if __name__ == "__main__":
    import asyncio
    asyncio.run(_example_upload())
